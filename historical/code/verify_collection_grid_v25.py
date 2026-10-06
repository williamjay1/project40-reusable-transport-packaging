"""Independent audits of collection/fate and country-generation component results.

Uses the independent source reconstruction, not either result producer. Confirms
scenario quantities and source component substitution; does not certify unknown
missing-asset treatment or national low-voltage electricity inventories.
"""
import argparse
import hashlib
import json
from pathlib import Path
import numpy as np
from openpyxl import load_workbook
from scipy.optimize import linprog
from verify_operating_thresholds_v24 import read_csv,reconstruct


def main():
    ap=argparse.ArgumentParser()
    for name in ['input-dir','stage-facts','collection-dir','grid-dir','grid-raw','adaptive-results','output']:
        ap.add_argument('--'+name,type=Path,required=True)
    a=ap.parse_args()
    polys=reconstruct(read_csv(a.input_dir/'tua_reported_grid.csv'),json.loads(a.stage_facts.read_text())['facts'],read_csv(a.input_dir/'tua_process_contributions.csv'))
    # Independent inversion checker polytopes are normalised by the printed U.
    gs=read_csv(a.input_dir/'tua_reported_grid.csv')
    units={r['indicator']:float(r['single_use_impact_per_service_100_crates']) for r in gs if r['source_scenario']=='W1_RE1'}
    results=json.loads((a.collection_dir/'calendar_network_results.json').read_text())['outputs']
    tab=read_csv(a.collection_dir/'collection_energy_fate_frontier.csv')
    rng=np.random.default_rng(400025); le=rng.triangular(1,1,2,50000); lh=rng.triangular(1,1,2,50000)
    maxerr=0.; checked=0
    for row in tab:
        r=next(r for r in results if r['success'] and r['mode']=='pooled' and r['scenario']['name']==row['scenario'])
        c=r['counts']; N=c['delivered_services']; W=c['wash_inputs']/100; L=c['missing_assets']; rho=float(row['missing_fate_virgin_production_equivalent_stress'])
        v=np.array([(c['virgin_assets']+rho*L)/100,c['terminal_recovered_assets']/100,c['closed_loop_rejections']/.55,c['closed_loop_rejections']/.55,W])
        tr=(c['empty_dirty_crate_km']+r['scenario']['clean_leg_factor']*c['empty_clean_crate_km'])/(139*100)-W
        estimates=[]; rectangle=[]
        for k,(A,b,u,p) in polys.items():
            lp=linprog(-v,A_ub=A,b_ub=b,bounds=[(None,None)]*5,method='highs'); assert lp.success
            base=float(v@lp.x)+max(tr*p['return_transport'])
            e=np.maximum((le-1)*W*p['electricity'][0],(le-1)*W*p['electricity'][1])
            h=np.maximum((lh-1)*W*p['water_heating'][0],(lh-1)*W*p['water_heating'][1])
            estimates.append((base+e+h)/(N/100*u))
            rectangle.append((base+W*max(p['electricity'])+W*max(p['water_heating']))/(N/100*u))
        mx=np.max(estimates,axis=0)
        vals=[float(np.mean(mx<1)),float(np.quantile(mx,.95)),float(max(rectangle))]
        keys=['design_prior_probability','design_prior_95pct_maximum_ratio','full_energy_rectangle_maximum_ratio']
        err=max(abs(v-float(row[k])) for k,v in zip(keys,vals)); maxerr=max(maxerr,err)
        assert err<1e-8,(row['scenario'],rho,err); checked+=len(keys)
    # Check inverse missing-fate allowances with the independently rebuilt support.
    allowances=read_csv(a.collection_dir/'missing_fate_allowances.csv')
    for row in allowances:
        r=next(r for r in results if r['scenario']['name']==row['scenario'] and r['mode']=='pooled')
        c=r['counts']; k=row['indicator']; A,b,u,p=polys[k]; W=c['wash_inputs']/100
        v=np.array([c['virgin_assets']/100,c['terminal_recovered_assets']/100,c['closed_loop_rejections']/.55,c['closed_loop_rejections']/.55,W])
        lp=linprog(-v,A_ub=A,b_ub=b,bounds=[(None,None)]*5,method='highs'); assert lp.success
        tr=(c['empty_dirty_crate_km']+r['scenario']['clean_leg_factor']*c['empty_clean_crate_km'])/(139*100)-W
        charge=float(v@lp.x)+max(tr*p['return_transport'])
        val=(c['delivered_services']/100*u-charge)*units[k]/c['missing_assets']
        assert abs(val-float(row['maximum_extra_missing_fate_burden_per_lost_asset']))<1e-8*max(abs(val),1e-20)
    # Direct raw-workbook extraction by indexed cell, a EU-only value check.
    manifest=json.loads((a.grid_dir/'generation_component_manifest.json').read_text())
    assert hashlib.sha256(a.grid_raw.read_bytes()).hexdigest()==manifest['source_sha256']
    wb=load_workbook(a.grid_raw,read_only=True,data_only=True); ws=wb['Table3_LC_GHG']
    year=[c.value for c in ws[3]].index(2024)+1
    raw={ws.cell(t,1).value:ws.cell(t,year).value for t in range(4,ws.max_row+1)}; wb.close()
    factors=read_csv(a.grid_dir/'eu27_generation_climate_factors.csv'); assert len(factors)==27
    for f in factors: assert float(f['generation_lifecycle_kgco2e_per_kwh'])==raw[f['geo']]
    original=json.loads(a.adaptive_results.read_text())['outputs']; countrytab=read_csv(a.grid_dir/'country_wash_climate_component.csv')
    A,b,u,p=polys['Climate change']; U=units['Climate change']
    for row in countrytab:
        r=next(r for r in original if r['scenario']['name']==row['scenario'] and r['mode']==row['mode'])
        c=r['counts']; N=c['delivered_services']; W=c['wash_inputs']/100; kwh=.0248*c['wash_inputs']
        v=np.array([c['virgin_assets']/100,c['terminal_recovered_assets']/100,c['closed_loop_rejections']/.55,c['closed_loop_rejections']/.55,W])
        lp=linprog(-v,A_ub=A,b_ub=b,bounds=[(None,None)]*5,method='highs'); assert lp.success
        tr=(c['empty_dirty_crate_km']+r['scenario']['clean_leg_factor']*c['empty_clean_crate_km'])/(139*100)-W
        upper=(float(v@lp.x)+max(tr*p['return_transport'])-W*min(p['electricity']))*U
        ef=float(row['generation_lifecycle_kgco2e_per_kwh']); target=N/100*u*U
        ratio=(upper+kwh*ef)/target; allowance=(target-upper-kwh*ef)/kwh
        assert abs(ratio-float(row['generation_only_climate_ratio_upper']))<1e-9
        assert abs(allowance-float(row['maximum_extra_delivery_climate_kgco2e_per_kwh']))<1e-8
    report=dict(all_passed=True,source='independent reconstructed source support, original facts and raw JRC workbook',
                collection_statistic_checks=checked,missing_fate_allowance_checks=len(allowances),
                eu27_raw_factor_checks=27,country_component_checks=len(countrytab)*2,
                maximum_collection_statistic_error=maxerr,
                scope='numerical verification, not empirical validation of missing fate, energy prior or electricity delivery')
    a.output.parent.mkdir(parents=True,exist_ok=True); a.output.write_text(json.dumps(report,indent=2)+'\n',encoding='utf-8')
    print(json.dumps(report,indent=2))


if __name__=='__main__': main()
