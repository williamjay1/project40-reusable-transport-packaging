"""EU-27 generation-only climate component sensitivity, not updated full LCIA.

JRC national generation factors lack a matched low-voltage delivery inventory.
The residual delivery intensity is therefore exposed as an inverse allowance.
Other impact categories and the source heat inventory are NOT multiplied by a
carbon ratio. Source print uncertainty is retained.
"""
import argparse
import hashlib
import json
from datetime import datetime, timezone
from pathlib import Path
from urllib.request import urlopen
import numpy as np
from openpyxl import load_workbook
from scipy.optimize import linprog
from calendar_pool_network_v23 import source_polytopes, dump_csv

URL='https://jeodpp.jrc.ec.europa.eu/ftp/jrc-opendata/CoM/EmissionsFactorNational2026/CoM-Emission-factors-for-national-electricity-generation-2026.xlsx'
EU={'AT','BE','BG','CY','CZ','DE','DK','EE','EL','ES','FI','FR','HR','HU','IE','IT','LT','LU','LV','MT','NL','PL','PT','RO','SE','SI','SK'}


def main():
    ap=argparse.ArgumentParser()
    for name in ['raw-file','output-dir','input-dir','stage-facts','network-results']:
        ap.add_argument('--'+name,type=Path,required=True)
    a=ap.parse_args(); out=a.output_dir; out.mkdir(parents=True,exist_ok=True)
    if not a.raw_file.exists():
        data=urlopen(URL,timeout=90).read(); assert data[:2]==b'PK' and len(data)<10000000
        a.raw_file.parent.mkdir(parents=True,exist_ok=True)
        with a.raw_file.open('xb') as f: f.write(data)
        a.raw_file.chmod(0o444)
    before=hashlib.sha256(a.raw_file.read_bytes()).hexdigest()
    wb=load_workbook(a.raw_file,read_only=True,data_only=True); ws=wb['Table3_LC_GHG']
    header=list(ws.values)[2]; column=header.index(2024)
    countries=[]
    for r in ws.iter_rows(min_row=4,values_only=True):
        if r[0] in EU:
            assert isinstance(r[column],(int,float)) and 0<=r[column]<2
            countries.append(dict(geo=r[0],country=r[1],year=2024,generation_lifecycle_kgco2e_per_kwh=float(r[column])))
    wb.close(); assert len(countries)==27 and len({r['geo'] for r in countries})==27
    assert before==hashlib.sha256(a.raw_file.read_bytes()).hexdigest()
    dump_csv(out/'eu27_generation_climate_factors.csv',countries)
    source=source_polytopes(a.input_dir,a.stage_facts); p=source['Climate change']; output=[]
    results=json.loads(a.network_results.read_text())['outputs']
    for r in results:
        if not r['success']: continue
        c=r['counts']; s=r['scenario']; N=c['delivered_services']; W=c['wash_inputs']/100
        v=np.array([c['virgin_assets']/100,c['terminal_recovered_assets']/100,c['closed_loop_rejections']/.55,c['closed_loop_rejections']/.55,W])
        sup=linprog(-v,A_ub=p['A'],b_ub=p['b'],bounds=[(None,None)]*5,method='highs')
        inf=linprog(v,A_ub=p['A'],b_ub=p['b'],bounds=[(None,None)]*5,method='highs')
        assert sup.success and inf.success
        t=(c['empty_dirty_crate_km']+s['clean_leg_factor']*c['empty_clean_crate_km'])/(139*100)-W
        tt=p['process']['return_transport']; te=p['process']['electricity']
        # Replacement removes the printed electricity component once; retain
        # its opposite endpoint when bounding the residual source burden.
        upper=(float(v@sup.x)+max(t*tt[0],t*tt[1])-W*te[0])*p['u']
        lower=(float(v@inf.x)+min(t*tt[0],t*tt[1])-W*te[1])*p['u']
        target=N/100*p['u']*p['u_lower_ratio']; kwh=.0248*c['wash_inputs']
        assert kwh>0
        for country in countries:
            ef=country['generation_lifecycle_kgco2e_per_kwh']
            charge=kwh*ef
            output.append(dict(scenario=s['name'],mode=r['mode'],**country,
                               source_climate_ratio_upper=r['impact_ratios']['Climate change']['upper'],
                               generation_only_climate_ratio_lower=(lower+charge)/(N/100*p['u']*p['u_upper_ratio']),
                               generation_only_climate_ratio_upper=(upper+charge)/target,
                               maximum_extra_delivery_climate_kgco2e_per_kwh=(target-upper-charge)/kwh,
                               wash_electricity_kwh=kwh,
                               scope='generation-only substitution with delivery gap exposed; not full national LCIA, not empirical relocation benefit'))
    dump_csv(out/'country_wash_climate_component.csv',output)
    manifest=dict(source_url=URL,landing_page='https://data.jrc.ec.europa.eu/dataset/0681adfd-9291-4b58-838a-67664b6d210f',
                  doi='10.2905/JRC.PWGK5H0',issued='2026-03-27',coverage='EU-27 2024 national generation LC GHG factors only',
                  accessed_at_utc=datetime.now(timezone.utc).isoformat(),source_sha256=before,license='CC BY 4.0; anonymous download',
                  extraction='Table3_LC_GHG; exact year column 2024; explicit EU27 whitelist; no non-EU factors or blank headers',
                  substitution='remove source electricity climate component once, add .0248 kWh/input * national generation factor; delivery gap remains unknown',
                  country_rows=len(countries),computed_rows=len(output),raw_unchanged=True)
    (out/'generation_component_manifest.json').write_text(json.dumps(manifest,indent=2)+'\n',encoding='utf-8')
    print(json.dumps(manifest,indent=2))


if __name__=='__main__': main()
